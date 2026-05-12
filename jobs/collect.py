"""
Phase 2 — COLLECT : scan data/collected/, complétude vs Phase 1, gold.retour_feuilles.
Chemins : data/collected/YYYY-MM-DD/<session>/ (noms identiques à generate.py).
"""
from __future__ import annotations

import argparse
import sys
from collections import defaultdict
from dataclasses import dataclass
from datetime import date, datetime
from decimal import Decimal
from pathlib import Path
from typing import Dict, List, Tuple

# Racine projet (parent de jobs/)
ROOT = Path(__file__).resolve().parent.parent
if str(ROOT) not in sys.path:
    sys.path.insert(0, str(ROOT))
if str(ROOT / "jobs") not in sys.path:
    sys.path.insert(0, str(ROOT / "jobs"))

if (ROOT / ".env").exists():
    try:
        from dotenv import load_dotenv

        load_dotenv(ROOT / ".env")
    except ImportError:
        pass

from openpyxl import load_workbook

import generate as gen
from utils.db import connect_postgres
from utils.logger import get_logger

logger = get_logger(__name__)

# Ligne d’en-tête colonnes dans generate.build_workbook (ligne 5) ; données à partir de la ligne 6
HEADER_ROW_GENERATE = 5


@dataclass
class StatFichier:
    cours: gen.CoursPlanifie
    nom_fichier: str
    statut: str  # "ok" | "manquant" | "vide" | "corrompu"
    detail: str


def parse_date(s: str | None) -> date:
    if not s:
        return date.today()
    return datetime.strptime(s, "%Y-%m-%d").date()


def valider_xlsx_collecte(path: Path) -> Tuple[bool, str]:
    """
    Retourne (True, "") si le fichier est lisible et contient au moins une ligne de données sous l’en-tête.
    """
    try:
        if not path.is_file():
            return False, "manquant"
        if path.stat().st_size == 0:
            return False, "vide (0 octet)"
        wb = load_workbook(path, read_only=True, data_only=True)
        try:
            ws = wb.active
            max_row = ws.max_row or 0
        finally:
            wb.close()
        if max_row <= HEADER_ROW_GENERATE:
            return False, "vide (aucune ligne étudiant sous l’en-tête)"
        return True, ""
    except Exception as exc:
        return False, f"corrompu ({exc})"


def attendus_pour_session(
    session: str,
    date_jour: date,
) -> List[Tuple[gen.CoursPlanifie, str]]:
    """Cours avec au moins un inscrit + nom de fichier = même règle que generate."""
    jour_label = gen.jour_semaine_fr(date_jour)
    cours_list = gen.fetch_cours_du_jour(session, jour_label)
    out: List[Tuple[gen.CoursPlanifie, str]] = []
    for c in cours_list:
        if not gen.fetch_inscrits(c.cours_id):
            continue
        out.append((c, gen.nom_fichier_xlsx(c)))
    return out


def collecter_feuilles(
    session: str,
    date_jour: date,
    base_collected: Path | None = None,
) -> Tuple[List[Path], List[StatFichier]]:
    """
    Compare collected vs attendu, met à jour gold.retour_feuilles.
    Retourne (chemins valides pour extract, détail par fichier).
    """
    base = base_collected or (ROOT / "data" / "collected")
    dossier = base / date_jour.isoformat() / session
    attendus = attendus_pour_session(session, date_jour)

    if not attendus:
        logger.info(
            "COLLECT : aucune feuille attendue (session=%s, date=%s)",
            session,
            date_jour.isoformat(),
        )
        return [], []

    # Clé en minuscules pour tolérer les différences de casse (Windows)
    collectes_ci: Dict[str, Path] = {}
    if dossier.is_dir():
        for p in dossier.glob("*.xlsx"):
            collectes_ci[p.name.lower()] = p

    stats: List[StatFichier] = []
    chemins_ok: List[Path] = []

    attendu_lower = {nom.lower() for _, nom in attendus}
    if dossier.is_dir():
        for p in dossier.glob("*.xlsx"):
            if p.name.lower() not in attendu_lower:
                logger.warning(
                    "Fichier collecté non attendu (ignoré pour le suivi) : %s",
                    p.name,
                )

    for c, fname in attendus:
        p = collectes_ci.get(fname.lower())
        if p is None:
            stats.append(
                StatFichier(
                    cours=c,
                    nom_fichier=fname,
                    statut="manquant",
                    detail="absent du dossier collected",
                )
            )
            logger.warning(
                "Feuille manquante : %s | enseignant %s %s | cours id=%s %s",
                fname,
                c.ens_prenom,
                c.ens_nom,
                c.cours_id,
                c.intitule,
            )
            continue

        ok, raison = valider_xlsx_collecte(p)
        if ok:
            stats.append(
                StatFichier(cours=c, nom_fichier=fname, statut="ok", detail="")
            )
            chemins_ok.append(p.resolve())
        else:
            st = "vide" if "vide" in raison else "corrompu"
            stats.append(
                StatFichier(cours=c, nom_fichier=fname, statut=st, detail=raison)
            )
            logger.warning(
                "Feuille invalide (%s) : %s | enseignant %s %s | cours id=%s",
                raison,
                fname,
                c.ens_prenom,
                c.ens_nom,
                c.cours_id,
            )

    # Agrégation par enseignant → gold.retour_feuilles
    by_ens: Dict[str, List[StatFichier]] = defaultdict(list)
    for s in stats:
        by_ens[s.cours.enseignant_id].append(s)

    with connect_postgres() as conn:
        with conn.cursor() as cur:
            for ens_id, lignes in sorted(by_ens.items()):
                nb_att = len(lignes)
                nb_ok = sum(1 for x in lignes if x.statut == "ok")
                taux = (Decimal(nb_ok) / Decimal(nb_att) * Decimal(100)).quantize(
                    Decimal("0.01")
                )
                cur.execute(
                    """
                    INSERT INTO gold.retour_feuilles (
                        enseignant_id, date_jour, session,
                        nb_attendus, nb_recus, taux_retour
                    )
                    VALUES (%s, %s, %s, %s, %s, %s)
                    ON CONFLICT (enseignant_id, date_jour, session) DO UPDATE SET
                        nb_attendus = EXCLUDED.nb_attendus,
                        nb_recus = EXCLUDED.nb_recus,
                        taux_retour = EXCLUDED.taux_retour,
                        computed_at = now()
                    """,
                    (ens_id, date_jour, session, nb_att, nb_ok, taux),
                )
        conn.commit()

    logger.info(
        "COLLECT terminé : %d fichier(s) valide(s) / %d attendu(s), dossier %s",
        len(chemins_ok),
        len(attendus),
        dossier,
    )
    return chemins_ok, stats


def main(argv: List[str] | None = None) -> int:
    p = argparse.ArgumentParser(
        description="Collecte : vérification des feuilles retournées et mise à jour gold.retour_feuilles."
    )
    p.add_argument(
        "--session",
        required=True,
        choices=gen.SESSIONS_VALIDES,
        help="matin ou apres-midi",
    )
    p.add_argument(
        "--date",
        default=None,
        help="Date au format YYYY-MM-DD (défaut : jour courant)",
    )
    args = p.parse_args(argv)

    date_jour = parse_date(args.date)
    try:
        collecter_feuilles(args.session, date_jour)
    except Exception:
        logger.exception("Échec COLLECT")
        return 1
    return 0


if __name__ == "__main__":
    sys.exit(main())
