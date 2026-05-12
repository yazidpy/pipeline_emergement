
import pandas as pd
from openpyxl import Workbook
import os

def create_test_excel(filename, data, header_row=7):
    wb = Workbook()
    ws = wb.active
    # On met des trucs avant
    for r in range(1, header_row):
        ws.cell(row=r, column=1, value=f"Header {r}")
    
    headers = ['Numéro Ét.', 'Nom', 'Prénom', 'Présent', 'Remarque']
    for c, h in enumerate(headers, start=1):
        ws.cell(row=header_row, column=c, value=h)
    
    for r, row_data in enumerate(data, start=header_row+1):
        for c, val in enumerate(row_data, start=1):
            ws.cell(row=r, column=c, value=val)
    
    wb.save(filename)
    print(f"File {filename} created.")

def validate_sim(filepath):
    print(f"\n--- Validating {filepath} ---")
    try:
        df = pd.read_excel(filepath, header=6)
        df.columns = ['Numéro Ét.', 'Nom', 'Prénom', 'Présent', 'Remarque']
        
        invalid_rows = []
        for idx, row in df.iterrows():
            val_p = str(row['Présent']).strip().upper() if pd.notnull(row['Présent']) else ""
            if val_p not in ['O', 'N']:
                stu_name = f"{row['Nom']} {row['Prénom']}"
                invalid_rows.append(f"Ligne {idx+8} : {stu_name} (reçu: '{val_p if val_p else 'VIDE'}')")
        
        if invalid_rows:
            print("ERROR Detected:")
            for err in invalid_rows:
                print(f"  {err}")
            return False
        else:
            print("SUCCESS: File is valid.")
            return True
    except Exception as e:
        print(f"CRITICAL ERROR: {e}")
        return False

# Data test sets
valid_data = [
    ["101", "Dupont", "Jean", "O", ""],
    ["102", "Martin", "Alice", "N", "Malade"]
]

invalid_data_empty = [
    ["101", "Dupont", "Jean", None, ""],
    ["102", "Martin", "Alice", "N", ""]
]

invalid_data_wrong = [
    ["101", "Dupont", "Jean", "O", ""],
    ["102", "Martin", "Alice", "X", "Quoi?"]
]

if __name__ == "__main__":
    os.makedirs("test_validation", exist_ok=True)
    
    create_test_excel("test_validation/valid.xlsx", valid_data)
    create_test_excel("test_validation/invalid_empty.xlsx", invalid_data_empty)
    create_test_excel("test_validation/invalid_wrong.xlsx", invalid_data_wrong)
    
    v1 = validate_sim("test_validation/valid.xlsx")
    v2 = validate_sim("test_validation/invalid_empty.xlsx")
    v3 = validate_sim("test_validation/invalid_wrong.xlsx")
    
    assert v1 == True
    assert v2 == False
    assert v3 == False
    
    print("\n✅ Verification SUCCESSful: All validation cases behaving as expected.")
